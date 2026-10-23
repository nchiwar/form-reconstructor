import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from pdf2docx import Converter


def export_to_excel(fields: list[dict], output_xlsx_path: str) -> str:
    """Creates a formatted Excel spreadsheet from the verified fields."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Extracted Form Data"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")

    headers = ["Page", "Field Label", "Verified Value", "Field Type"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for item in fields:
        ws.append([
            item.get("page_num", 0) + 1,
            item.get("field_label", ""),
            item.get("handwritten_value", ""),
            item.get("field_type", "text"),
        ])

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 5, 15)

    wb.save(output_xlsx_path)
    return output_xlsx_path


def append_to_master_template_excel(template_id: str, fields: list[dict], templates_dir: str) -> str:
    """
    Appends a new submission as a single horizontal row in a Master Excel file
    for organizations processing the same form repeatedly.
    """
    master_xlsx = os.path.join(templates_dir, f"{template_id}_master.xlsx")
    text_fields = [f for f in fields if f.get("field_type") != "signature"]

    if not os.path.exists(master_xlsx):
        wb = Workbook()
        ws = wb.active
        ws.title = "Master Submissions"
        headers = ["Timestamp"] + [f.get("field_label", f"Field {i+1}") for i, f in enumerate(text_fields)]
        ws.append(headers)
    else:
        wb = load_workbook(master_xlsx)
        ws = wb.active

    row_values = [datetime.now().strftime("%Y-%m-%d %H:%M:%S")] + [
        f.get("handwritten_value", "") for f in text_fields
    ]
    ws.append(row_values)
    wb.save(master_xlsx)
    return master_xlsx


def export_to_word(clean_pdf_path: str, fields: list[dict], output_docx_path: str) -> str:
    """
    Converts the reconstructed PDF into an editable Word (.docx) document.
    Falls back to a structured Word table if vector conversion fails.
    """
    try:
        cv = Converter(clean_pdf_path)
        cv.convert(output_docx_path, start=0, end=None)
        cv.close()
        return output_docx_path
    except Exception:
        doc = Document()
        doc.add_heading("Reconstructed Form Data", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Page"
        hdr_cells[1].text = "Field Label"
        hdr_cells[2].text = "Verified Value"

        for item in fields:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item.get("page_num", 0) + 1)
            row_cells[1].text = str(item.get("field_label", ""))
            row_cells[2].text = str(item.get("handwritten_value", ""))

        doc.save(output_docx_path)
        return output_docx_path